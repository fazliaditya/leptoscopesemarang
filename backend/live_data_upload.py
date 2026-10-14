"""Upload, validation, merge, and template helpers for LeptoScope case data.

Designed to work without extra Python packages beyond the current LeptoScope
requirements. CSV is parsed with pandas. XLSX is parsed with openpyxl when it is
available and otherwise with a small stdlib XLSX reader for simple tabular files.
Uploads are sent as raw request bodies (application/octet-stream), so
python-multipart is not required.
"""
from __future__ import annotations

from datetime import datetime, timedelta, timezone
from io import BytesIO, StringIO
from pathlib import Path
from urllib.parse import quote
from zipfile import ZipFile, ZIP_DEFLATED
import json
import re
import xml.etree.ElementTree as ET

import numpy as np
import pandas as pd

JAKARTA_TZ = timezone(timedelta(hours=7))
MONTH_NAMES = {
    1: "Januari", 2: "Februari", 3: "Maret", 4: "April", 5: "Mei", 6: "Juni",
    7: "Juli", 8: "Agustus", 9: "September", 10: "Oktober", 11: "November", 12: "Desember",
}


def normalize_code(value):
    return str(value or "").strip().replace(" ", "")


def normalize_name(value):
    return " ".join(str(value or "").lower().strip().split())


def month_label(value) -> str:
    ts = pd.Timestamp(value)
    return f"{MONTH_NAMES[int(ts.month)]} {int(ts.year)}"


def _month_start(now: datetime) -> pd.Timestamp:
    return pd.Timestamp(now.year, now.month, 1)


def required_periods(now: datetime | None = None):
    now = now or datetime.now(JAKARTA_TZ)
    current = _month_start(now)
    completed = [current - pd.DateOffset(months=offset) for offset in range(4, 0, -1)]
    return {
        "current": current,
        "completed": completed,
        "input_month": completed[-1],
        "target_month": current,
    }


def district_reference(model_df: pd.DataFrame):
    rows = (
        model_df[["kecamatan_code", "kecamatan_name"]]
        .dropna()
        .drop_duplicates()
        .copy()
    )
    rows["kecamatan_code"] = rows["kecamatan_code"].map(normalize_code)
    rows["kecamatan_name"] = rows["kecamatan_name"].astype(str).str.strip()
    return rows.sort_values("kecamatan_code").reset_index(drop=True)


def _aliases(frame: pd.DataFrame):
    aliases = {
        "date": ["date", "tanggal", "bulan", "month", "periode"],
        "code": ["kecamatan_code", "kode_kecamatan", "district_code", "kode"],
        "name": ["kecamatan_name", "nama_kecamatan", "kecamatan", "district_name"],
        "count": ["case_count", "jumlah_kasus", "kasus", "cases", "jumlah kasus"],
    }
    lowered = {str(column).strip().lower(): column for column in frame.columns}
    return {
        key: next((lowered[candidate] for candidate in candidates if candidate in lowered), None)
        for key, candidates in aliases.items()
    }


def _col_index(cell_ref: str) -> int:
    letters = re.match(r"[A-Z]+", cell_ref or "A")
    text = letters.group(0) if letters else "A"
    result = 0
    for char in text:
        result = result * 26 + (ord(char) - 64)
    return result - 1


def _read_xlsx_stdlib(content: bytes) -> pd.DataFrame:
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    rel_ns = {"r": "http://schemas.openxmlformats.org/package/2006/relationships"}
    office_rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    with ZipFile(BytesIO(content)) as archive:
        shared = []
        if "xl/sharedStrings.xml" in archive.namelist():
            root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
            for si in root.findall("m:si", ns):
                shared.append("".join(node.text or "" for node in si.findall(".//m:t", ns)))

        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        first_sheet = workbook.find("m:sheets/m:sheet", ns)
        if first_sheet is None:
            raise ValueError("Workbook tidak memiliki sheet.")
        rel_id = first_sheet.attrib.get(f"{{{office_rel}}}id")
        rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        target = None
        for rel in rels.findall("r:Relationship", rel_ns):
            if rel.attrib.get("Id") == rel_id:
                target = rel.attrib.get("Target")
                break
        if not target:
            raise ValueError("Sheet pertama tidak dapat dibaca.")
        if target.startswith("/"):
            sheet_path = target.lstrip("/")
        else:
            sheet_path = "xl/" + target.lstrip("/")
        sheet_path = re.sub(r"xl/\.\./", "", sheet_path)
        sheet = ET.fromstring(archive.read(sheet_path))
        rows = []
        max_col = 0
        for row in sheet.findall("m:sheetData/m:row", ns):
            values = {}
            for cell in row.findall("m:c", ns):
                idx = _col_index(cell.attrib.get("r", "A1"))
                max_col = max(max_col, idx)
                cell_type = cell.attrib.get("t")
                value_node = cell.find("m:v", ns)
                inline_node = cell.find("m:is", ns)
                value = ""
                if cell_type == "inlineStr" and inline_node is not None:
                    value = "".join(node.text or "" for node in inline_node.findall(".//m:t", ns))
                elif value_node is not None:
                    raw = value_node.text or ""
                    if cell_type == "s":
                        try:
                            value = shared[int(raw)]
                        except Exception:
                            value = raw
                    else:
                        value = raw
                values[idx] = value
            if values:
                rows.append([values.get(i, "") for i in range(max_col + 1)])
        if not rows:
            return pd.DataFrame()
        width = max(len(row) for row in rows)
        rows = [row + [""] * (width - len(row)) for row in rows]
        headers = [str(value).strip() for value in rows[0]]
        return pd.DataFrame(rows[1:], columns=headers)


def read_uploaded_table(filename: str, content: bytes) -> pd.DataFrame:
    if not content:
        raise ValueError("File kosong.")
    suffix = Path(filename or "").suffix.lower()
    if suffix == ".csv":
        text = None
        for encoding in ("utf-8-sig", "utf-8", "cp1252"):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise ValueError("Encoding CSV tidak dapat dibaca. Gunakan CSV UTF-8.")
        return pd.read_csv(StringIO(text), dtype=str, keep_default_na=False)
    if suffix == ".xlsx":
        try:
            import openpyxl  # type: ignore
            workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
            sheet = workbook[workbook.sheetnames[0]]
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                return pd.DataFrame()
            headers = [str(value or "").strip() for value in rows[0]]
            return pd.DataFrame(rows[1:], columns=headers)
        except ModuleNotFoundError:
            return _read_xlsx_stdlib(content)
    raise ValueError("Format file harus .xlsx atau .csv.")


def _parse_date_value(value):
    if value is None or str(value).strip() == "":
        return pd.NaT
    if isinstance(value, (datetime, pd.Timestamp)):
        return pd.Timestamp(value)
    text = str(value).strip()
    # Excel serial dates can appear when using the lightweight XLSX parser.
    try:
        serial = float(text)
        if serial > 20000:
            return pd.Timestamp("1899-12-30") + pd.to_timedelta(serial, unit="D")
    except Exception:
        pass
    return pd.to_datetime(text, errors="coerce")


def _normalize_frame(frame: pd.DataFrame, reference: pd.DataFrame, current_month: pd.Timestamp):
    columns = _aliases(frame)
    errors = []
    warnings = []
    if columns["date"] is None:
        errors.append("Kolom tanggal/bulan tidak ditemukan.")
    if columns["count"] is None:
        errors.append("Kolom case_count/jumlah_kasus tidak ditemukan.")
    if columns["code"] is None and columns["name"] is None:
        errors.append("Kolom kode atau nama kecamatan tidak ditemukan.")
    if errors:
        return pd.DataFrame(), errors, warnings, []

    code_to_name = dict(zip(reference["kecamatan_code"], reference["kecamatan_name"]))
    name_to_code = {normalize_name(name): code for code, name in code_to_name.items()}
    normalized = []
    preview = []

    for index, row in frame.iterrows():
        raw_date = row.get(columns["date"])
        date = _parse_date_value(raw_date)
        raw_code = row.get(columns["code"]) if columns["code"] else ""
        raw_name = row.get(columns["name"]) if columns["name"] else ""
        code = normalize_code(raw_code)
        name_key = normalize_name(raw_name)
        canonical_code = code if code in code_to_name else name_to_code.get(name_key, "")
        canonical_name = code_to_name.get(canonical_code, "")
        raw_count = row.get(columns["count"])
        blank_count = raw_count is None or str(raw_count).strip() == ""
        status = "Valid"
        issue = ""

        if pd.isna(date):
            status = "Tidak valid"; issue = "Tanggal/bulan tidak valid"
        elif not canonical_code:
            status = "Tidak valid"; issue = "Kecamatan tidak dikenali"
        elif pd.Timestamp(date.year, date.month, 1) > current_month:
            status = "Tidak valid"; issue = "Periode berada setelah bulan berjalan"
        elif blank_count:
            status = "Kosong"; issue = "Belum ada nilai kasus"
        else:
            try:
                count = float(str(raw_count).replace(",", "."))
                if not np.isfinite(count) or count < 0:
                    raise ValueError
                if not np.isclose(count, round(count)):
                    status = "Tidak valid"; issue = "Jumlah kasus harus bilangan bulat"
                else:
                    count = int(round(count))
            except Exception:
                status = "Tidak valid"; issue = "Jumlah kasus tidak valid"

        preview.append({
            "row": int(index) + 2,
            "period": "—" if pd.isna(date) else pd.Timestamp(date.year, date.month, 1).strftime("%Y-%m"),
            "kecamatan": canonical_name or str(raw_name or raw_code or "—"),
            "case_count": None if blank_count else raw_count,
            "status": status,
            "issue": issue,
        })
        if status == "Tidak valid":
            errors.append(f"Baris {int(index)+2}: {issue}.")
            continue
        if blank_count:
            continue
        month = pd.Timestamp(date.year, date.month, 1)
        normalized.append({
            "date": month,
            "kecamatan_code": canonical_code,
            "kecamatan_name": canonical_name,
            "case_count": int(count),
        })

    clean = pd.DataFrame(normalized)
    if not clean.empty:
        duplicated = clean.duplicated(["kecamatan_code", "date"], keep=False)
        if duplicated.any():
            duplicate_rows = clean.loc[duplicated, ["kecamatan_name", "date"]]
            pairs = ", ".join(
                f"{row.kecamatan_name} {month_label(row.date)}"
                for row in duplicate_rows.itertuples(index=False)
            )
            errors.append(f"Duplikasi kecamatan-bulan pada file: {pairs}.")
    return clean, list(dict.fromkeys(errors)), warnings, preview


def _load_active(live_dir: Path, reference: pd.DataFrame, current_month: pd.Timestamp):
    path = Path(live_dir) / "kasus_terbaru.csv"
    if not path.exists():
        return pd.DataFrame(columns=["date", "kecamatan_code", "kecamatan_name", "case_count"])
    try:
        frame = pd.read_csv(path, dtype=str, keep_default_na=False)
        clean, _, _, _ = _normalize_frame(frame, reference, current_month)
        return clean
    except Exception:
        return pd.DataFrame(columns=["date", "kecamatan_code", "kecamatan_name", "case_count"])


def _merge_active(active: pd.DataFrame, uploaded: pd.DataFrame):
    if active.empty:
        merged = uploaded.copy()
    elif uploaded.empty:
        merged = active.copy()
    else:
        active = active.copy(); active["source_order"] = 0
        uploaded = uploaded.copy(); uploaded["source_order"] = 1
        merged = pd.concat([active, uploaded], ignore_index=True)
        merged = merged.sort_values("source_order").drop_duplicates(["kecamatan_code", "date"], keep="last")
        merged = merged.drop(columns=["source_order"])
    if merged.empty:
        return merged
    return merged.sort_values(["date", "kecamatan_code"]).reset_index(drop=True)


def summarize_candidate(candidate: pd.DataFrame, reference: pd.DataFrame, now: datetime | None = None):
    now = now or datetime.now(JAKARTA_TZ)
    periods = required_periods(now)
    completed = periods["completed"]
    current = periods["current"]
    required_slots = len(reference) * len(completed)
    lookup = {}
    if not candidate.empty:
        for row in candidate.itertuples(index=False):
            lookup[(normalize_code(row.kecamatan_code), pd.Timestamp(row.date))] = int(row.case_count)

    missing = []
    valid_slots = 0
    for month in completed:
        for district in reference.itertuples(index=False):
            key = (normalize_code(district.kecamatan_code), pd.Timestamp(month))
            if key in lookup:
                valid_slots += 1
            else:
                missing.append({
                    "period": pd.Timestamp(month).strftime("%Y-%m"),
                    "period_label": month_label(month),
                    "kecamatan": district.kecamatan_name,
                })

    current_rows = candidate[candidate["date"] == current].copy() if not candidate.empty else pd.DataFrame()
    current_total = int(current_rows["case_count"].sum()) if not current_rows.empty else 0
    current_districts = int(current_rows["kecamatan_code"].nunique()) if not current_rows.empty else 0
    latest_period = candidate["date"].max() if not candidate.empty else None
    return {
        "model_ready": valid_slots == required_slots,
        "model_valid_slots": valid_slots,
        "model_required_slots": required_slots,
        "model_missing_slots": required_slots - valid_slots,
        "missing_preview": missing[:12],
        "completed_periods": [pd.Timestamp(month).strftime("%Y-%m") for month in completed],
        "completed_period_labels": [month_label(month) for month in completed],
        "model_input_period": periods["input_month"].strftime("%Y-%m"),
        "model_input_label": month_label(periods["input_month"]),
        "prediction_target": current.strftime("%Y-%m"),
        "prediction_target_label": month_label(current),
        "monitoring_period": current.strftime("%Y-%m"),
        "monitoring_period_label": month_label(current),
        "monitoring_districts_reported": current_districts,
        "monitoring_total_cases": current_total,
        "monitoring_complete": current_districts == len(reference),
        "latest_period": latest_period.strftime("%Y-%m") if latest_period is not None else None,
    }


def validate_upload(filename: str, content: bytes, model_df: pd.DataFrame, live_dir: Path, now: datetime | None = None):
    now = now or datetime.now(JAKARTA_TZ)
    reference = district_reference(model_df)
    periods = required_periods(now)
    frame = read_uploaded_table(filename, content)
    if frame.empty:
        raise ValueError("File tidak memiliki baris data.")
    uploaded, errors, warnings, preview = _normalize_frame(frame, reference, periods["current"])
    active = _load_active(Path(live_dir), reference, periods["current"])
    candidate = _merge_active(active, uploaded)
    summary = summarize_candidate(candidate, reference, now)
    valid_uploaded_rows = int(len(uploaded))
    if valid_uploaded_rows == 0 and not errors:
        warnings.append("Tidak ada nilai kasus yang terisi pada file. Sel kosong dianggap data belum diterima.")
    return {
        "filename": filename,
        "can_apply": not errors and valid_uploaded_rows > 0,
        "errors": errors[:20],
        "warnings": warnings[:20],
        "uploaded_valid_rows": valid_uploaded_rows,
        "uploaded_preview": preview[:24],
        "summary": summary,
    }, candidate


def apply_upload(filename: str, content: bytes, model_df: pd.DataFrame, live_dir: Path, now: datetime | None = None):
    now = now or datetime.now(JAKARTA_TZ)
    result, candidate = validate_upload(filename, content, model_df, live_dir, now)
    if not result["can_apply"]:
        raise ValueError("Data belum dapat diterapkan. Periksa hasil validasi.")
    live_dir = Path(live_dir)
    live_dir.mkdir(parents=True, exist_ok=True)
    save = candidate.copy()
    save["tanggal"] = pd.to_datetime(save["date"]).dt.strftime("%Y-%m-01")
    save = save[["tanggal", "kecamatan_code", "kecamatan_name", "case_count"]]
    path = live_dir / "kasus_terbaru.csv"
    save.to_csv(path, index=False, encoding="utf-8-sig")
    meta = {
        "applied_at": now.isoformat(),
        "source_filename": filename,
        "rows_active": int(len(save)),
        "model_ready": result["summary"]["model_ready"],
        "monitoring_period": result["summary"]["monitoring_period"],
        "monitoring_total_cases": result["summary"]["monitoring_total_cases"],
        "monitoring_districts_reported": result["summary"]["monitoring_districts_reported"],
    }
    (live_dir / "case_upload_meta.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
    result["applied"] = True
    result["applied_at"] = now.isoformat()
    result["active_rows"] = int(len(save))
    return result


def live_status(model_df: pd.DataFrame, live_dir: Path, now: datetime | None = None):
    now = now or datetime.now(JAKARTA_TZ)
    reference = district_reference(model_df)
    periods = required_periods(now)
    active = _load_active(Path(live_dir), reference, periods["current"])
    summary = summarize_candidate(active, reference, now)
    meta_path = Path(live_dir) / "case_upload_meta.json"
    upload_meta = {}
    if meta_path.exists():
        try:
            upload_meta = json.loads(meta_path.read_text(encoding="utf-8"))
        except Exception:
            upload_meta = {}
    return {
        **summary,
        "active_rows": int(len(active)),
        "has_active_file": (Path(live_dir) / "kasus_terbaru.csv").exists(),
        "applied_at": upload_meta.get("applied_at"),
        "source_filename": upload_meta.get("source_filename"),
        "district_count": int(len(reference)),
        "note": (
            "Data bulan berjalan dapat diperbarui kapan saja untuk monitoring. "
            "Model peringatan tetap memakai empat bulan kalender lengkap terakhir."
        ),
    }


def _xml_escape(value):
    text = str(value if value is not None else "")
    return (text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            .replace('"', "&quot;").replace("'", "&apos;"))


def _excel_col(index: int):
    index += 1
    letters = ""
    while index:
        index, rem = divmod(index - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


def _worksheet_xml(rows):
    xml_rows = []
    for r_index, row in enumerate(rows, start=1):
        cells = []
        for c_index, value in enumerate(row):
            ref = f"{_excel_col(c_index)}{r_index}"
            if isinstance(value, (int, float)) and value is not None:
                cells.append(f'<c r="{ref}"><v>{value}</v></c>')
            else:
                cells.append(f'<c r="{ref}" t="inlineStr"><is><t>{_xml_escape(value)}</t></is></c>')
        xml_rows.append(f'<row r="{r_index}">' + "".join(cells) + "</row>")
    return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            '<sheetData>' + "".join(xml_rows) + '</sheetData></worksheet>')


def build_template_xlsx(model_df: pd.DataFrame, now: datetime | None = None) -> bytes:
    now = now or datetime.now(JAKARTA_TZ)
    reference = district_reference(model_df)
    periods = required_periods(now)
    months = periods["completed"] + [periods["current"]]
    data_rows = [["tanggal", "kecamatan_code", "kecamatan_name", "case_count", "status_periode"]]
    for month in months:
        status = "Bulan berjalan - monitoring sementara" if month == periods["current"] else "Bulan lengkap - input model"
        for district in reference.itertuples(index=False):
            data_rows.append([
                pd.Timestamp(month).strftime("%Y-%m-01"),
                district.kecamatan_code,
                district.kecamatan_name,
                "",
                status,
            ])
    guide_rows = [
        ["Template Pembaruan Data Kasus LeptoScope"],
        ["Periode target", month_label(periods["current"])],
        ["Input model", ", ".join(month_label(month) for month in periods["completed"])],
        ["Aturan 1", "Isi case_count dengan bilangan bulat >= 0."],
        ["Aturan 2", "0 berarti nihil kasus terverifikasi; kosong berarti data belum diterima."],
        ["Aturan 3", "Empat bulan lengkap terakhir harus memiliki 16 kecamatan agar dipakai model."],
        ["Aturan 4", "Bulan berjalan boleh parsial dan hanya digunakan untuk monitoring terkini."],
        ["Aturan 5", "Upload ulang akan memperbarui kecamatan-bulan yang sama, bukan menjumlahkan ulang."],
    ]
    content_types = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>
<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>
<Override PartName="/xl/worksheets/sheet2.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>
</Types>'''
    root_rels = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>
</Relationships>'''
    workbook = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
<sheets><sheet name="Data Kasus" sheetId="1" r:id="rId1"/><sheet name="Petunjuk" sheetId="2" r:id="rId2"/></sheets></workbook>'''
    workbook_rels = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>
<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet2.xml"/>
</Relationships>'''
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", root_rels)
        archive.writestr("xl/workbook.xml", workbook)
        archive.writestr("xl/_rels/workbook.xml.rels", workbook_rels)
        archive.writestr("xl/worksheets/sheet1.xml", _worksheet_xml(data_rows))
        archive.writestr("xl/worksheets/sheet2.xml", _worksheet_xml(guide_rows))
    return output.getvalue()


def template_filename(now: datetime | None = None):
    periods = required_periods(now or datetime.now(JAKARTA_TZ))
    return f"Template_Kasus_LeptoScope_{periods['current'].strftime('%Y_%m')}.xlsx"
